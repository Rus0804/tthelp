"""
excel_viewer.py  –  Streamlit app to display Excel files preserving
merged cells, cell colours, fonts, borders, and number formats.

Run:
    streamlit run excel_viewer.py
"""

import io
import re
import streamlit as st
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.styles.colors import COLOR_INDEX, aRGB_REGEX
import os, pathlib
from typing import Optional

INDEXED_COLOURS = [
    "000000","FFFFFF","FF0000","00FF00","0000FF","FFFF00","FF00FF","00FFFF",
    "000000","FFFFFF","FF0000","00FF00","0000FF","FFFF00","FF00FF","00FFFF",
    "800000","008000","000080","808000","800080","008080","C0C0C0","808080",
    "9999FF","993366","FFFFCC","CCFFFF","660066","FF8080","0066CC","CCCCFF",
    "000080","FF00FF","FFFF00","00FFFF","800080","800000","008080","0000FF",
    "00CCFF","CCFFFF","CCFFCC","FFFF99","99CCFF","FF99CC","CC99FF","FFCC99",
    "3366FF","33CCCC","99CC00","FFCC00","FF9900","FF6600","666699","969696",
    "003366","339966","003300","333300","993300","993366","333399","333333",
]

def resolve_colour(colour_obj) -> str | None:
    """Return a CSS hex colour string or None."""
    if colour_obj is None:
        return "#FFFFFF"
    t = getattr(colour_obj, "type", None)
    val = getattr(colour_obj, "rgb", None) or getattr(colour_obj, "value", None)
    if t == "rgb" and val and val not in ("00000000", "FFFFFFFF", "FF000000"):
        rgb = val[-6:] if len(val) == 8 else val
        if rgb != "000000" or t == "rgb":
            return f"#{rgb}"
    if t == "indexed":
        idx = getattr(colour_obj, "indexed", None)
        if idx is not None and 0 <= idx < len(INDEXED_COLOURS):
            return f"#{INDEXED_COLOURS[idx]}"
    return None

def bg_from_fill(fill) -> str | None:
    if fill is None or fill.fill_type in (None, "none"):
        return None
    fg = getattr(fill, "fgColor", None)
    return resolve_colour(fg)

def colour_from_font(font) -> str | None:
    if font is None:
        return None
    return resolve_colour(getattr(font, "color", None))


def border_style_css(side) -> str:
    if side is None or side.border_style is None:
        return "none"
    mapping = {
        "thin":        "1px solid #999",
        "medium":      "2px solid #666",
        "thick":       "3px solid #333",
        "dashed":      "1px dashed #999",
        "dotted":      "1px dotted #999",
        "double":      "3px double #999",
        "hair":        "1px solid #ccc",
        "mediumDashed":"2px dashed #666",
    }
    return mapping.get(side.border_style, "1px solid #ccc")


def fmt_value(cell) -> str:
    val = cell.value
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        nf = cell.number_format or ""
        if "%" in nf:
            return f"{val*100:.1f}%"
        if any(c in nf for c in ("#,##", "0,0")):
            return f"{val:,.0f}"
        if "$" in nf:
            return f"${val:,.2f}"
    return str(val)

def build_merge_map(ws):
    """
    Returns:
      skip  – set of (row,col) cells that are covered by a merge and should be skipped
      spans – dict (top-left row,col) → (rowspan, colspan)
    """
    skip, spans = set(), {}
    for rng in ws.merged_cells.ranges:
        r1, c1, r2, c2 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        spans[(r1, c1)] = (r2 - r1 + 1, c2 - c1 + 1)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if (r, c) != (r1, c1):
                    skip.add((r, c))
    return skip, spans

# ─────────────────────────── HTML renderer ────────────────────────────────

def sheet_to_html(ws) -> str:
    skip, spans = build_merge_map(ws)

    # column widths (approximate: 1 Excel unit ≈ 7 px)
    col_widths = {}
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        dim = ws.column_dimensions.get(letter)
        col_widths[col_idx] = int((dim.width if dim and dim.width else 8) * 7)

    rows_html = []
    for row in ws.iter_rows():
        # row height
        rh = ws.row_dimensions.get(row[0].row)
        h_px = int((rh.height if rh and rh.height else 15) * 1.33)
        cells_html = []
        for cell in row:
            r, c = cell.row, cell.column
            if (r, c) in skip:
                continue

            rs, cs = spans.get((r, c), (1, 1))
            rs_attr = f' rowspan="{rs}"' if rs > 1 else ""
            cs_attr = f' colspan="{cs}"' if cs > 1 else ""

            # ── styles ──
            styles = []
            styles.append(f"width:{col_widths.get(c, 56)}px")
            styles.append(f"height:{h_px}px")
            styles.append("padding:2px 4px")
            styles.append("overflow:hidden")
            styles.append("white-space:pre-wrap")
            styles.append("word-break:break-word")

            bg = bg_from_fill(cell.fill)
            if bg:
                styles.append(f"background:{bg}")
            else:
                styles.append(f"background:#FFFFFF")

            font = cell.font
            if font:
                if font.bold:
                    styles.append("font-weight:bold")
                if font.italic:
                    styles.append("font-style:italic")
                if font.underline:
                    styles.append("text-decoration:underline")
                if font.size:
                    styles.append(f"font-size:{int(font.size)}px")
                fc = colour_from_font(font)
                if fc:
                    styles.append(f"color:{fc}")

            align = cell.alignment
            if align:
                ha = align.horizontal or "left"
                va = align.vertical or "top"
                va_map = {"center": "middle", "top": "top", "bottom": "bottom"}
                styles.append(f"text-align:{ha}")
                styles.append(f"vertical-align:{va_map.get(va, 'top')}")
            else:
                styles.append("text-align:left")
                styles.append("vertical-align:top")

            bd = cell.border
            if bd:
                styles.append(f"border-top:{border_style_css(bd.top)}")
                styles.append(f"border-right:{border_style_css(bd.right)}")
                styles.append(f"border-bottom:{border_style_css(bd.bottom)}")
                styles.append(f"border-left:{border_style_css(bd.left)}")
            else:
                styles.append("border:1px solid #e0e0e0")

            style_str = ";".join(styles)
            content = fmt_value(cell)

            cells_html.append(
                f'<td{rs_attr}{cs_attr} style="{style_str}">{content}</td>'
            )

        rows_html.append(f"<tr>{''.join(cells_html)}</tr>")

    table = (
        '<html><body style="margin:0;padding:0;background:#ffffff">'  # ← wrap in html/body
        '<div style="overflow:auto;border:1px solid #ddd;border-radius:4px;width:max-content">'
        '<table style="border-collapse:collapse;font-family:Calibri,Arial,sans-serif;font-size:13px">'
        + "".join(rows_html)
        + "</table></div>"
        '</body></html>'
    )
    return table


def display_default(file: Optional[UploadedFile] = None):
    if not file:
        file_path = "TERM 4 MBA TT.xlsx"
        with open(file_path, "rb") as f:
            wb = load_workbook(f, data_only=True)
    else:
        wb = load_workbook(file, data_only = True)
    ws = wb.active
    html = sheet_to_html(ws)
    st.iframe(html, height=700)