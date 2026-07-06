from utils import get_row_range, build_merge_map
import openpyxl
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
import pandas as pd
import os
import streamlit as st

def get_course_dict(file: Optional[UploadedFile] = None):

    if file:
        wb = load_workbook(file)
    else:
        wb = load_workbook("TERM 4 MBA TT.xlsx")
    ws = wb.active

    row_start, _ = get_row_range(ws['A'], 'codes')

    courses_dict = {}

    merge_map     = build_merge_map(ws)
    already_seen  = set()
    cleared       = 0
    
    for row in range(row_start, len(ws['A'])+ 1):
        for col in range(1, 11):
            cell = ws.cell(row=row, column=col)

            if isinstance(cell, MergedCell):
                master_coords = merge_map.get((row, col))
                if not master_coords or master_coords in already_seen:
                    continue
                already_seen.add(master_coords)
                master = ws.cell(*master_coords)
            else:
                if (row, col) in already_seen:
                    continue
                already_seen.add((row, col))
                master = cell

            if col == 1:
                course_code = master.value[:8].strip()
                instructors = master.value[9:].strip()
                courses_dict[course_code] = {"Instructor": instructors, "Course Name": '', "Credits": 0, "Faculty": [], "Students": '', "Venue": []}
            elif col == 4:
                courses_dict[course_code]["Course Name"] = master.value
            elif col == 6:
                courses_dict[course_code]["Credits"] = int(master.value)
            elif col == 7:
                segments = master.value.split('(', 1)
                name = segments[0].strip()
                segments.pop(0)

                if(len(segments)==0):
                    classes = []
                elif len(segments[-1].split(',')) > 1:
                    classes = segments[-1].split(',')
                    classes[-1] = classes[-1].strip()[:-1]
                else:
                    classes = segments[-1].split('&')
                    classes[-1] = classes[-1].strip()[:-1]
                
                for i in range(len(classes)):
                    classes[i] = classes[i].strip()

                faculty_dict = {"Name": name, "Sections": classes}
                courses_dict[course_code]["Faculty"].append(faculty_dict)
            elif col == 9:
                courses_dict[course_code]["Students"] = master.value
            elif col == 10:
                courses_dict[course_code]["Venue"].append(master.value)
    
    return courses_dict
